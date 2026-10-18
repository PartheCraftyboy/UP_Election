#!/usr/bin/env python3
"""Link the photos in mla_photo/ into deliverables.

  python3 link_photos.py              -> xlsx + html (html uses relative paths)
  python3 link_photos.py --embed      -> html with photos inlined as base64
                                         (one self-contained file, no folder needed)

Matching is by the leading 3-digit constituency number, so you can rename the
rest of the file however you like (005_Brijesh.jpg, 005_whatever.png - all fine)
and .jpg/.jpeg/.png/.webp all work. Any row with no matching file just renders
blank instead of breaking.
"""
import base64, csv, glob, html, io, os, re, sys
from PIL import Image
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

SRC, DIR, SIZE = "uttar_pradesh_mlas_2026.csv", "mla_photo", 50
EMBED = "--embed" in sys.argv

# index every image by its leading number -> tolerant of renamed files
index = {}
for p in glob.glob(os.path.join(DIR, "*")):
    m = re.match(r"(\d{1,3})", os.path.basename(p))
    if m and p.lower().endswith((".png", ".jpg", ".jpeg", ".webp")):
        index.setdefault(int(m.group(1)), p)

rows = list(csv.DictReader(open(SRC, encoding="utf-8")))
cols = [c for c in rows[0].keys() if c != "Photo_File"]

def square(path):
    """Centre-crop to a square and resize to 50x50. Returns a PNG buffer."""
    im = Image.open(path).convert("RGB")
    w, h = im.size
    s = min(w, h)
    im = im.crop(((w - s) // 2, (h - s) // 2, (w + s) // 2, (h + s) // 2))
    im = im.resize((SIZE, SIZE), Image.LANCZOS)
    buf = io.BytesIO(); im.save(buf, "PNG"); buf.seek(0)
    return buf

# ---------- XLSX with real embedded images ----------
wb = Workbook(); ws = wb.active; ws.title = "UP MLAs 2026"
hdr = PatternFill("solid", fgColor="1F3864")
ws.append(["Photo"] + cols)
for c in range(1, len(cols) + 2):
    cell = ws.cell(1, c); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hdr
ws.freeze_panes = "B2"
ws.column_dimensions["A"].width = 9

missing = 0
for i, r in enumerate(rows, start=2):
    ws.append([None] + [r[c] for c in cols])
    ws.row_dimensions[i].height = 40
    p = index.get(int(r["Constituency_No"]))
    if p:
        img = XLImage(square(p)); img.width = img.height = SIZE
        img.anchor = f"A{i}"
        ws.add_image(img)
    else:
        missing += 1
    ws.cell(i, 1).alignment = Alignment(horizontal="center", vertical="center")

for j, c in enumerate(cols, start=2):
    width = max(len(c), *(len(str(r[c])) for r in rows)) + 2
    ws.column_dimensions[get_column_letter(j)].width = min(width, 42)
wb.save("uttar_pradesh_mlas_2026.xlsx")

# ---------- HTML ----------
def src(p):
    if not p:
        return ""
    if EMBED:
        return "data:image/png;base64," + base64.b64encode(square(p).read()).decode()
    return p.replace(os.sep, "/")

cells = []
for r in rows:
    p = index.get(int(r["Constituency_No"]))
    img = f'<img src="{src(p)}" width="50" height="50" alt="">' if p else '<span class="np"></span>'
    cells.append("<tr><td class=\"ph\">" + img + "</td>" +
                 "".join(f"<td>{html.escape(str(r[c]))}</td>" for c in cols) + "</tr>")

doc = f"""<!doctype html><meta charset="utf-8"><title>UP MLAs 2026</title>
<style>
 body{{font:14px/1.45 system-ui,sans-serif;margin:24px;color:#1a1a1a}}
 h1{{font-size:20px;margin:0 0 4px}} p.sub{{color:#666;margin:0 0 18px}}
 table{{border-collapse:collapse;font-size:13px}}
 th{{background:#1f3864;color:#fff;padding:8px 10px;text-align:left;
     position:sticky;top:0}}
 td{{border-bottom:1px solid #e6e6e6;padding:6px 10px;vertical-align:middle}}
 td.ph{{width:58px;padding:4px 6px}}
 img{{border-radius:6px;display:block}}
 .np{{display:block;width:50px;height:50px;border-radius:6px;background:#eee}}
 tr:hover td{{background:#f7f9fc}}
</style>
<h1>18th Uttar Pradesh Vidhan Sabha &mdash; 403 seats</h1>
<p class="sub">Ratings and economic values are estimates, not official figures.</p>
<table><thead><tr><th>Photo</th>{''.join(f'<th>{html.escape(c)}</th>' for c in cols)}</tr></thead>
<tbody>{''.join(cells)}</tbody></table>"""
open("uttar_pradesh_mlas_2026.html", "w", encoding="utf-8").write(doc)

print(f"linked {len(rows)-missing}/{len(rows)} photos; {missing} missing")
